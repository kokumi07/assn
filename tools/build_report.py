"""Build Korean PDF/DOCX documentation from the captured, actual game runs.

Documentation dependencies: python-docx, playwright, Chromium and Noto CJK fonts.
The game and its tests do not depend on these packages.
"""

import base64
import html
from pathlib import Path
import shutil
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from playwright.sync_api import sync_playwright


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
SHOTS = DOCS / "screenshots"
FONT = "Noto Sans CJK KR"
CHOICE_PROMPT = "Rock (1), Paper (2), Scissors (3), Quit (0): "
LEGEND = "○-Player ●-Computer ◐-Player & Computer"


def paragraph(text):
    return {"kind": "p", "text": text}


def heading(text):
    return {"kind": "h2", "text": text}


def figure(name, caption, width=174):
    return {"kind": "figure", "name": name, "caption": caption, "width": width}


def table(headers, rows):
    return {"kind": "table", "headers": headers, "rows": rows}


def read_capture(name):
    return (DOCS / "transcripts" / f"{name}.txt").read_text(encoding="utf-8")


def round_excerpt(text, number):
    """Crop a continuous section, from choices through the next position line."""
    starts = []
    offset = 0
    while True:
        index = text.find("Player: ", offset)
        if index < 0:
            break
        starts.append(index)
        offset = index + 1
    start = starts[number]
    position = text.index("Player = ", start)
    end = text.find("\n", position)
    return text[start:end if end >= 0 else len(text)].rstrip()


def last_board(text):
    start = text.index("[", text.rindex(LEGEND))
    return text[start:].rstrip()


def take_screenshots(browser):
    mixed = read_capture("mixed")
    captures = {
        "input": read_capture("validation").split(CHOICE_PROMPT)[0].rstrip(),
        "quit": read_capture("quit")[read_capture("quit").index(CHOICE_PROMPT):].rstrip(),
        "round_player": round_excerpt(mixed, 0),
        "round_computer": round_excerpt(mixed, 1),
        "round_tie": round_excerpt(mixed, 2),
        "final_player": last_board(read_capture("player_win")),
        "final_computer": last_board(read_capture("computer_win")),
    }
    page = browser.new_page(viewport={"width": 700, "height": 1200}, device_scale_factor=2)
    for name, text in captures.items():
        # Text is never retyped or altered: these are continuous crops of PTY logs.
        content = (
            '<!doctype html><html lang="ko"><meta charset="utf-8"><style>'
            'html,body{margin:0;background:#fff;}pre{box-sizing:border-box;'
            'width:700px;margin:0;padding:10px 14px;border:1px solid #cbd5e1;'
            'border-radius:5px;background:#f7f9fc;color:#111827;'
            'font:14px/19px "Noto Sans Mono CJK KR",monospace;white-space:pre;}'
            '</style><pre>' + html.escape(text) + '</pre></html>'
        )
        page.set_content(content)
        page.evaluate("document.fonts.ready")
        element = page.locator("pre")
        if element.evaluate("e => e.scrollWidth > e.clientWidth"):
            raise RuntimeError(f"Clipped screenshot text: {name}")
        element.screenshot(path=str(SHOTS / f"{name}.png"))
    page.close()


def draw_structure(browser):
    svg = '''<svg xmlns="http://www.w3.org/2000/svg" width="960" height="300" viewBox="0 0 960 300">
    <rect width="960" height="300" fill="white"/>
    <g fill="none" stroke="#8ea9c4" stroke-width="2.5">
      <path d="M480 76 V110 M160 110 H800 M160 110 V142 M480 110 V142 M800 110 V142"/>
    </g>
    <rect x="330" y="12" width="300" height="64" rx="9" fill="#173e65"/>
    <text x="480" y="40" text-anchor="middle" fill="white" font-size="23">play_game()</text>
    <text x="480" y="63" text-anchor="middle" fill="white" font-size="17">입력·라운드 반복·종료 관리</text>
    <g fill="#eff5fb" stroke="#9fb9d0" stroke-width="1.5">
      <rect x="12" y="142" width="296" height="138" rx="8"/>
      <rect x="332" y="142" width="296" height="138" rx="8"/>
      <rect x="652" y="142" width="296" height="138" rx="8"/>
    </g>
    <g text-anchor="middle" fill="#173e65" font-size="22">
      <text x="160" y="174">입력 / 선택</text><text x="480" y="174">판정 / 이동</text><text x="800" y="174">출력</text>
    </g>
    <g text-anchor="middle" fill="#243648" font-size="20">
      <text x="160" y="211">get_player_choice()</text><text x="160" y="245">get_computer_choice()</text>
      <text x="480" y="211">determine_winner(...)</text><text x="480" y="245">update_position(...)</text>
      <text x="800" y="211">print_stairs(...)</text><text x="800" y="245" font-size="17">계단 + 현재 위치</text>
    </g></svg>'''
    (DOCS / "structure.svg").write_text(svg, encoding="utf-8")
    page = browser.new_page(viewport={"width": 960, "height": 300}, device_scale_factor=2)
    page.set_content('<html><body style="margin:0;font-family:Noto Sans CJK KR">' + svg + '</body></html>')
    page.evaluate("document.fonts.ready")
    page.locator("svg").screenshot(path=str(SHOTS / "structure.png"))
    page.close()


def report_pages():
    return [
        {"title": "", "cover": True, "blocks": []},
        {"title": "1. 프로그램 개요와 구조", "blocks": [
            paragraph("플레이어와 컴퓨터가 같은 높이에서 시작하여 가위바위보의 승자만 계단을 내려가는 콘솔 게임이다. 먼저 0단계에 도착한 쪽이 최종 승리한다. 계단 수는 5~20이며, 매 라운드의 선택·승패·현재 위치를 화면에 표시한다."),
            heading("1.1 이동 규칙과 표시"),
            table(["승리한 선택", "입력값", "이동"], [
                ["Rock (바위)", "1", "1단계 아래로"],
                ["Paper (보)", "2", "2단계 아래로"],
                ["Scissors (가위)", "3", "3단계 아래로"],
                ["무승부", "두 선택이 같음", "양쪽 위치 유지"],
            ]),
            paragraph("승자는 자신이 고른 숫자만큼 내려간다. 이동 후 위치가 음수가 되면 0으로 보정한다. ○는 플레이어, ●는 컴퓨터, ◐는 같은 위치의 두 참가자를 뜻한다. 계단은 위에서부터 0까지 출력하며, s단계의 ▨ 개수는 ‘전체 계단 수 − s’이다."),
            heading("1.2 함수 구조도"),
            figure("structure", "그림 1. play_game()이 필수 함수 5개를 호출하는 전체 구조", 174),
            paragraph("게임 상태인 전체 계단 수, 두 위치, 현재 선택과 판정 결과는 함수의 지역 변수로 관리한다. 필요한 값은 매개변수와 반환값으로 전달하며, 전역 변수와 추가 사용자 정의 함수는 사용하지 않았다. 계단 수 입력, 종료 확인, 안내 문구 출력은 play_game()이 담당한다."),
        ]},
        {"title": "2. 함수와 알고리즘", "blocks": [
            paragraph("필수 함수 6개의 이름과 매개변수는 과제 명세 그대로 유지했다. 반환값은 다음과 같이 정했다. 반환값이 없는 출력·제어 함수는 None을 반환한다."),
            table(["함수 / 매개변수", "역할 / 반환값"], [
                ["get_player_choice()\n매개변수 없음", "0~3 정수 입력을 반복 검증\n반환: 유효한 정수"],
                ["get_computer_choice()\n매개변수 없음", "random.randint(1, 3)으로 선택\n반환: 1, 2, 3 중 정수"],
                ["print_stairs(total_steps,\nplayer_pos, computer_pos)", "전체 계단과 두 위치 출력\n반환: None"],
                ["determine_winner(player_choice,\ncomputer_choice)", "두 선택을 비교하여 승패 결정\n반환: 'player', 'computer', 'tie'"],
                ["update_position(current_pos,\nmove_steps)", "이동 후 위치를 계산하고 0 하한 보정\n반환: 보정된 위치 (정수)"],
                ["play_game()\n매개변수 없음", "입력부터 최종 종료까지 진행 제어\n반환: None"],
            ]),
            heading("2.1 전체 의사코드"),
            {"kind": "pre", "text": "① 제목 출력 → 5~20의 정수가 입력될 때까지 재입력\n② 두 위치를 전체 계단 수로 초기화 → 계단 출력\n③ 어느 쪽도 0이 아닌 동안 다음을 반복\n   · 0~3의 유효한 플레이어 선택을 입력받음\n   · 선택이 0이면 yes/no를 유효할 때까지 확인\n     yes: Game Over!를 출력하고 종료 / no: ③으로 복귀\n   · 컴퓨터 선택 생성 → 승패 판정 → 선택과 결과 출력\n   · ENTER 입력 대기\n   · 무승부이면 유지, 아니면 승자 위치에서 승자 선택값 차감\n     결과가 음수이면 0으로 보정\n   · 계단과 현재 위치를 다시 출력\n④ 0에 먼저 도착한 참가자의 최종 승리 메시지 출력"},
            heading("2.2 판정과 종료 조건"),
            paragraph("같은 선택이면 무승부이다. 플레이어의 승리 조합은 (1, 3), (2, 1), (3, 2)이며, 나머지는 컴퓨터 승리이다. 게임이 끝나는 라운드에도 ENTER 입력과 마지막 계단 출력을 수행한 후 최종 승자를 알린다. 시작 지점에서는 __main__ 조건을 통해 play_game()을 한 번 호출한다."),
        ]},
        {"title": "3. 실행 방법과 입력 검증", "blocks": [
            paragraph("Python 3.x가 설치된 터미널에서 assn1.py가 있는 폴더로 이동한 뒤 ‘python assn1.py’를 실행한다. 별도 패키지나 데이터 파일은 필요 없다. UTF-8 계단 기호를 지원하는 글꼴을 사용한다. 검증 환경은 Linux, Python 3.12.14이다."),
            paragraph("계단 수는 5~20, 게임 선택은 1(바위)·2(보)·3(가위)·0(종료)이다. 라운드 결과 뒤 ENTER를 누르면 이동한다. 과제의 가정에 따라 숫자 입력은 정수로 받으며, 정수가 아닌 입력에 대한 예외 처리는 추가하지 않았다."),
            figure("input", "그림 2. 3과 123은 거부하고 5를 허용한다. 두 참가자는 5단계에서 함께 시작한다."),
            figure("quit", "그림 3. NO로 종료를 취소한 후 0을 다시 입력하면 재확인한다. yse는 거부하고 YES는 허용한다."),
            {"kind": "note", "text": "그림 2~8은 별도 Python 프로세스의 실제 PTY 출력에서 연속 구간을 발췌해 캡처했다. 재현을 위해 캡처 도구에서만 random.randint를 정해진 선택 순서로 대체했다. 일반 실행에서는 매 라운드 무작위로 선택하므로 같은 입력의 결과가 달라질 수 있다."},
        ]},
        {"title": "4. 라운드 실행 예시", "blocks": [
            paragraph("전체 계단 수를 5로 정하고 컴퓨터 선택을 3 → 2 → 2로 고정한 한 실행의 앞 두 라운드이다. 두 경우 모두 결과를 표시한 다음 ENTER 입력을 받고 위치를 갱신한다."),
            figure("round_player", "그림 4. 플레이어 1(바위), 컴퓨터 3(가위): 플레이어만 5 → 4로 1단계 이동한다."),
            figure("round_computer", "그림 5. 플레이어 1(바위), 컴퓨터 2(보): 컴퓨터만 5 → 3으로 2단계 이동한다."),
            paragraph("그림 4의 새 위치 (4, 5)가 그림 5의 시작 상태이다. 컴퓨터가 이겼을 때 플레이어의 선택값 1이 아니라 컴퓨터의 선택값 2를 이동량으로 사용한다. 이처럼 승패 판정과 위치 계산을 분리해 양쪽에 동일한 이동 함수를 적용한다."),
        ]},
        {"title": "5. 무승부와 최종 종료 예시", "blocks": [
            figure("round_tie", "그림 6. 그림 5 다음 라운드에서 양쪽이 보를 선택하면 무승부이다. 위치 (4, 3)가 유지된다."),
            figure("final_player", "그림 7. 별도 실행: 플레이어가 가위로 두 번 승리해 5 → 2 → 0으로 이동한다. 2 − 3은 0으로 보정된다."),
            figure("final_computer", "그림 8. 별도 실행: 컴퓨터가 보·보·가위로 승리해 5 → 3 → 1 → 0으로 이동한 뒤 종료한다."),
            paragraph("그림 7과 8은 마지막 계단부터 종료 메시지까지의 연속 구간이다. 양측 모두 0보다 아래로 표시되지 않으며, 최종 위치를 출력한 뒤 추가 라운드 입력 없이 종료된다."),
        ]},
        {"title": "6. 검증 및 고찰", "blocks": [
            paragraph("‘python -m unittest discover -s tests -v’로 23개 테스트를 실행해 모두 통과했다. 9가지 승패 조합 등은 하위 사례로 검사하며 테스트 메서드 수와 구분한다. 일반 CLI의 실제 시작·종료도 별도 프로세스로 확인했다."),
            table(["검증 항목", "확인한 조건", "결과"], [
                ["승패 / 난수", "선택 9조합, randint(1, 3) 호출", "통과"],
                ["계단 수", "5·20 허용, 4·21·음수 거부", "통과"],
                ["선택 입력", "0~3 허용, −1·4·123 재입력", "통과"],
                ["계단 표시", "초기·같은 위치·분리·0·두 자리 정렬", "통과"],
                ["이동 / 순서", "승자 선택값, ENTER 후 갱신, 0 보정", "통과"],
                ["게임 흐름", "무승부 유지, 양측 최종 승리", "통과"],
                ["종료 확인", "대소문자, 오타, 취소 후 재확인", "통과"],
                ["독립 실행", "CLI 종료 코드 0, import 시 자동 실행 없음", "통과"],
            ]),
            heading("6.1 구현상의 검토와 결정"),
            paragraph("① 같은 위치에서는 두 기호를 붙이지 않고 ◐ 하나를 우선 출력했다. ② 승자 선택값을 이동량으로 전달해 컴퓨터 승리에서도 올바른 거리를 이동하도록 했다. ③ 종료 입력과 잘못된 선택은 라운드로 처리하지 않으므로 컴퓨터 난수를 뽑지 않는다. ④ 결과 → ENTER → 이동 → 재출력 순서를 지켜 마지막 라운드도 동일하게 처리했다."),
            paragraph("검증에서 발견된 실행 실패나 미해결 기능 결함은 없었다. 정수가 아닌 숫자 입력, 입력 스트림의 강제 종료, 글꼴 미지원은 명세 밖의 상황으로 처리하지 않았다. 개인의 개발 경험을 추정하는 대신 코드와 검증에서 확인한 결정 사항을 정리했다."),
            heading("6.2 결론"),
            paragraph("본 구현은 입력 검증, 승패 판정, 위치 계산과 출력을 함수로 나누고 지역 상태와 반환값으로 연결한다. 이 구조는 승패 9조합과 이동 경계처럼 작은 단위의 검증을 가능하게 하며, 전체 게임에서는 반복문과 종료 조건의 순서를 명확히 보여 준다."),
            heading("6.3 향후 개선"),
            paragraph("별도 확장에서는 비정수 입력과 EOF 처리, 기호를 지원하지 않는 터미널의 ASCII 표시, 재시작 메뉴를 고려할 수 있다. 현재 버전에는 명세의 출력 형식을 유지하기 위해 이러한 추가 기능을 넣지 않았다."),
            heading("참고 자료"),
            {"kind": "small", "text": "[1] assn1_2026f.pdf, CSED101 Programming Assignment #1 (기능 명세).\n[2] AssnReadMe.pdf (보고서 구성과 파일 형식).\n[3] report_ex1.pdf, report_ex2.pdf (보고서 구성·시각 형식 참고).\n실행 근거: docs/transcripts/ · 검증 코드: tests/test_assn1.py"},
        ]},
    ]


CSS = '''
@page { size: A4; margin: 18mm; }
* { box-sizing: border-box; }
body { margin: 0; color: #172a3a; font-family: "Noto Sans CJK KR", sans-serif; font-size: 10.2pt; line-height: 1.57; }
.page { height: 260mm; position: relative; break-after: page; padding: 0 0 11mm; }
.page:last-child { break-after: auto; }
.running { font-size: 8pt; color: #637587; border-bottom: 1px solid #d5dee7; padding-bottom: 3mm; margin-bottom: 5mm; display:flex; justify-content:space-between; }
h1 { font-size: 19pt; line-height: 1.3; margin: 0 0 5mm; letter-spacing: -.3pt; color: #173e65; }
h2 { font-size: 11.2pt; margin: 4mm 0 2mm; color: #173e65; }
p { margin: 0 0 3mm; word-break: keep-all; overflow-wrap: anywhere; }
table { width:100%; border-collapse:collapse; margin:3mm 0; font-size:9.2pt; line-height:1.4; }
th { background:#eaf1f8; color:#173e65; text-align:left; font-weight:700; }
th,td { padding:2.1mm 2.5mm; border:1px solid #cbd6e1; vertical-align:top; }
pre { margin:2mm 0 3mm; padding:3mm; background:#f2f6fa; border-left:3px solid #8eacc8; white-space:pre-wrap; font:9pt/1.58 "Noto Sans CJK KR",sans-serif; }
figure { margin:2.2mm 0 3.5mm; break-inside:avoid; }
figure img { max-width:100%; height:auto; display:block; margin:auto; }
figcaption { font-size:8.5pt; line-height:1.45; color:#43596c; margin-top:1.5mm; }
.note { font-size:8.2pt; color:#526577; border-left:2px solid #9fb9d0; padding-left:2.5mm; }
.small { font-size:8.1pt; line-height:1.5; }
.footer { position:absolute; bottom:0; left:0; right:0; font-size:8pt; color:#637587; border-top:1px solid #d5dee7; padding-top:2mm; display:flex; justify-content:space-between; }
.cover { padding-top:18mm; }
.cover .eyebrow { font-size:10pt; color:#617b94; letter-spacing:1pt; margin-bottom:18mm; }
.cover h1 { font-size:31pt; line-height:1.4; margin-bottom:5mm; }
.cover .subtitle { font-size:15pt; color:#4e657c; }
.cover .rule { width:22mm; height:1mm; background:#173e65; margin:12mm 0; }
.cover table { width:115mm; font-size:11pt; margin-top:16mm; }
.cover th,.cover td { padding:3mm; background:transparent; border:0; border-bottom:1px solid #e0e6ed; }
.cover th { width:32mm; }
.disclosure { margin-top:15mm; padding:4mm; background:#f2f6fa; font-size:9.5pt; }
.disclosure strong { display:block; margin-bottom:1.5mm; }
'''


def block_html(block):
    kind = block["kind"]
    escape = lambda text: html.escape(text).replace("\n", "<br>")
    if kind in ("p", "h2", "pre"):
        return f'<{kind}>' + html.escape(block["text"]) + f'</{kind}>'
    if kind in ("note", "small"):
        return f'<p class="{kind}">' + escape(block["text"]) + '</p>'
    if kind == "table":
        header = ''.join('<th>' + escape(cell) + '</th>' for cell in block["headers"])
        rows = ''.join('<tr>' + ''.join('<td>' + escape(cell) + '</td>' for cell in row) + '</tr>' for row in block["rows"])
        return '<table><thead><tr>' + header + '</tr></thead><tbody>' + rows + '</tbody></table>'
    if kind == "figure":
        return f'<figure><img style="width:{block["width"]}mm" src="screenshots/{block["name"]}.png"><figcaption>' + escape(block["caption"]) + '</figcaption></figure>'
    raise ValueError(kind)


def write_html(pages):
    content = ['<!doctype html><html lang="ko"><head><meta charset="utf-8"><title>CSED101 Assignment #1 — 고현우</title><style>' + CSS + '</style></head><body>']
    for number, page in enumerate(pages, 1):
        footer = f'<div class="footer"><span>CSED101 · Assignment #1 · 고현우</span><span>{number} / {len(pages)}</span></div>'
        if page.get("cover"):
            content.append('''<section class="page cover">
              <div class="eyebrow">CSED101 · FALL 2026</div>
              <h1>Programming &amp;<br>Problem Solving</h1>
              <div class="subtitle">Assignment #1</div><div class="rule"></div>
              <h2 style="font-size:21pt">Rock-Paper-Scissors<br>Stair Game</h2>
              <p style="color:#617b94">가위바위보 계단 게임 · 프로그램 보고서</p>
              <table><tr><th>학과</th><td>무은재학부</td></tr><tr><th>학번</th><td>20260134</td></tr>
              <tr><th>이름</th><td>고현우</td></tr><tr><th>POVIS ID</th><td>kohyunwoo</td></tr></table>
              <div class="disclosure"><strong>작성 방식</strong>코드 구현, 보고서 작성 및 검증에 생성형 AI의 도움을 받았다.</div>
            ''' + footer + '</section>')
        else:
            content.append('<section class="page"><div class="running"><span>Programming &amp; Problem Solving</span><span>2026 Fall · ASSN 1</span></div>')
            content.append('<h1>' + html.escape(page["title"]) + '</h1><main>')
            content.extend(block_html(block) for block in page["blocks"])
            content.append('</main>' + footer + '</section>')
    content.append('</body></html>')
    destination = DOCS / "report.html"
    destination.write_text('\n'.join(content), encoding="utf-8")
    return destination


def set_font(run, size=None, bold=None, color=None):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_text(document, text, size=None, bold=False, color=None, after=5):
    p = document.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.2
    set_font(p.add_run(text), size=size, bold=bold, color=color)
    return p


def keep_table_row(row):
    properties = row._tr.get_or_add_trPr()
    properties.append(OxmlElement('w:cantSplit'))


def write_docx(pages):
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin = section.bottom_margin = Mm(18)
    section.left_margin = section.right_margin = Mm(18)
    section.header_distance = section.footer_distance = Mm(9)
    normal = document.styles['Normal']
    normal.font.name = FONT
    normal.font.size = Pt(10)
    normal._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), FONT)
    normal.paragraph_format.line_spacing = 1.2
    normal.paragraph_format.space_after = Pt(5)
    for name, size in [('Title', 29), ('Heading 1', 18), ('Heading 2', 11)]:
        style = document.styles[name]
        style.font.name = FONT
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string('173E65')
        style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), FONT)
        style.paragraph_format.space_before = Pt(9)
        style.paragraph_format.space_after = Pt(5)
    header = section.header.paragraphs[0]
    set_font(header.add_run('CSED101  ·  Programming & Problem Solving  ·  Fall 2026'), 8, color='637587')
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_font(footer.add_run('고현우 · 20260134    |    '), 8, color='637587')
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)

    for index, page in enumerate(pages):
        if index:
            document.add_page_break()
        if page.get('cover'):
            add_text(document, 'CSED101 · FALL 2026', 11, color='637587', after=48)
            add_text(document, 'Programming &\nProblem Solving', 30, True, '173E65', 10)
            add_text(document, 'Assignment #1', 17, color='526B83', after=32)
            add_text(document, 'Rock-Paper-Scissors\nStair Game', 23, True, '173E65', 8)
            add_text(document, '가위바위보 계단 게임 · 프로그램 보고서', 12, color='526B83', after=42)
            for label, value in [('학과', '무은재학부'), ('학번', '20260134'), ('이름', '고현우'), ('POVIS ID', 'kohyunwoo')]:
                add_text(document, f'{label}     {value}', 12, after=13)
            add_text(document, '작성 방식', 11, True, '173E65', after=5).paragraph_format.space_before = Pt(30)
            add_text(document, '코드 구현, 보고서 작성 및 검증에 생성형 AI의 도움을 받았다.', 10)
            continue
        document.add_heading(page['title'], level=1)
        for block in page['blocks']:
            kind = block['kind']
            if kind == 'h2':
                document.add_heading(block['text'], level=2)
            elif kind in ('p', 'note', 'small', 'pre'):
                p = add_text(document, block['text'], size=8.3 if kind in ('note', 'small') else 9 if kind == 'pre' else 10)
                if kind in ('note', 'small'):
                    for run in p.runs:
                        set_font(run, color='526577')
                if kind == 'pre':
                    p.paragraph_format.line_spacing = 1.15
                    p.paragraph_format.keep_together = True
            elif kind == 'table':
                t = document.add_table(rows=1, cols=len(block['headers']))
                t.style = 'Table Grid'
                t.alignment = WD_TABLE_ALIGNMENT.CENTER
                for cell, text in zip(t.rows[0].cells, block['headers']):
                    cell.text = text
                    shading = OxmlElement('w:shd')
                    shading.set(qn('w:fill'), 'EAF1F8')
                    cell._tc.get_or_add_tcPr().append(shading)
                    for run in cell.paragraphs[0].runs:
                        set_font(run, 9, True, '173E65')
                header_repeat = OxmlElement('w:tblHeader')
                t.rows[0]._tr.get_or_add_trPr().append(header_repeat)
                for row_data in block['rows']:
                    row = t.add_row()
                    keep_table_row(row)
                    for cell, text in zip(row.cells, row_data):
                        cell.text = text
                        for p in cell.paragraphs:
                            p.paragraph_format.space_after = Pt(3)
                            p.paragraph_format.space_before = Pt(3)
                            p.paragraph_format.line_spacing = 1.08
                            for run in p.runs:
                                set_font(run, 8.8)
                document.add_paragraph().paragraph_format.space_after = Pt(0)
            elif kind == 'figure':
                p = document.add_paragraph()
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.keep_with_next = True
                p.add_run().add_picture(str(SHOTS / f'{block["name"]}.png'), width=Mm(block['width']))
                p = add_text(document, block['caption'], 8.3, color='43596C', after=7)
                p.paragraph_format.keep_together = True
            else:
                raise ValueError(kind)
    document.core_properties.title = 'CSED101 Assignment #1 — Rock-Paper-Scissors Stair Game'
    document.core_properties.author = '고현우'
    document.core_properties.subject = '무은재학부 · 20260134 · kohyunwoo'
    document.save(ROOT / 'assn1.docx')


def write_markdown(pages):
    parts = ['# CSED101 — Assignment #1', '', 'Rock-Paper-Scissors Stair Game', '',
             '무은재학부 · 20260134 · 고현우 · kohyunwoo', '',
             '작성 방식: 코드 구현, 보고서 작성 및 검증에 생성형 AI의 도움을 받았다.', '']
    for page in pages:
        if page.get('cover'):
            continue
        parts.extend(['## ' + page['title'], ''])
        for b in page['blocks']:
            kind = b['kind']
            if kind == 'h2':
                parts.append('### ' + b['text'])
            elif kind == 'figure':
                parts.extend([f'![{b["caption"]}](screenshots/{b["name"]}.png)', b['caption']])
            elif kind == 'table':
                parts.append('| ' + ' | '.join(b['headers']) + ' |')
                parts.append('| ' + ' | '.join('---' for _ in b['headers']) + ' |')
                parts.extend('| ' + ' | '.join(c.replace('\n', '<br>') for c in row) + ' |' for row in b['rows'])
            elif kind == 'pre':
                parts.extend(['```text', b['text'], '```'])
            else:
                parts.append(b['text'])
            parts.append('')
    (DOCS / 'report.md').write_text('\n'.join(parts), encoding='utf-8')


def main():
    SHOTS.mkdir(parents=True, exist_ok=True)
    chromium = shutil.which('chromium') or shutil.which('chromium-browser')
    if not chromium:
        raise SystemExit('Chromium is required only to regenerate the report.')
    with sync_playwright() as p:
        browser = p.chromium.launch(executable_path=chromium, headless=True, args=['--no-sandbox'])
        take_screenshots(browser)
        draw_structure(browser)
        pages = report_pages()
        source = write_html(pages)
        page = browser.new_page()
        # Inline authorized local image bytes; the managed browser disables file://.
        rendered_html = source.read_text(encoding='utf-8')
        for image_path in SHOTS.glob('*.png'):
            encoded = base64.b64encode(image_path.read_bytes()).decode('ascii')
            rendered_html = rendered_html.replace(
                f'src="screenshots/{image_path.name}"', f'src="data:image/png;base64,{encoded}"')
        page.set_content(rendered_html, wait_until='load')
        page.evaluate('document.fonts.ready')
        overflow = page.evaluate('''() => [...document.querySelectorAll('.page')].map((s,i) => {
          const content = s.querySelector('main');
          return content && content.getBoundingClientRect().bottom > s.querySelector('.footer').getBoundingClientRect().top - 8 ? i+1 : null;
        }).filter(Boolean)''')
        if overflow:
            raise RuntimeError(f'Report content overlaps footer on pages: {overflow}')
        page.pdf(path=str(ROOT / 'assn1.pdf'), format='A4', print_background=True, prefer_css_page_size=True,
                 display_header_footer=False)
        browser.close()
    write_docx(pages)
    write_markdown(pages)
    print('Created assn1.pdf, assn1.docx, docs/report.html and docs/report.md')


if __name__ == '__main__':
    main()
